"""Turning one published document into passages worth quoting.

Adapted from the customer support corpus reader, narrowed to the formats a
ministry, a commission or an agency actually issues. Ma'at answers a claim by
quoting a sentence back and highlighting it, so a passage has to be prose
somebody signed, carrying enough of its position to be cited: which page it sat
on, which heading it sat under.

Two things are load-bearing and easy to lose in a rewrite.

**Position travels with the text.** A PDF is read page by page rather than as
one string, so every passage knows its page. Headings become a path, so a
fragment from deep inside a circular still says what it was under. Without
that, a citation can name a document but never a place inside it, and a
visitor cannot check the quote.

**A file that cannot be read says so.** Every reader below opens a container
written by somebody else, and corrupt or truncated input makes those libraries
raise their own exception types. Left alone they surface as a 500 that tells
the uploader nothing. One boundary turns all of them into the same actionable
refusal.
"""

from __future__ import annotations

import io
import re

from llama_index.core import Document as LlamaDocument
from llama_index.core.node_parser import MarkdownNodeParser, SentenceSplitter

#: Roughly a long paragraph. Big enough to carry an argument, small enough that
#: a highlighted sentence is easy to find inside the quote shown to a visitor.
CHUNK_TOKENS = 512
CHUNK_OVERLAP = 64

#: Formats the upload door takes. Mirrored in the widget's dashboard; the
#: server is what actually gates.
#:
#: No .doc: the legacy binary format needs LibreOffice on the host to convert,
#: which is a service dependency for a format bodies have largely stopped
#: issuing. No .html: a stored page is a stored script.
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".odt", ".rtf", ".txt")

#: Characters that print as nothing and break everything: zero-width spaces
#: and joiners, word joiners, soft hyphens, byte-order marks. Some publishers'
#: editors sprinkle them through prose; left in, "fees" is no longer "fees" to
#: full-text search and the embedding sees a word it has never met.
ZERO_WIDTH = re.compile("[\u200b\u200c\u200d\u200e\u200f\u2060\u2061\u2062\u2063\u2064\ufeff\u00ad]")


#: The tell-tale of text that was UTF-8 once and got read as Windows-1252 on
#: the way into a publisher's own system: "â€™" for an apostrophe, "â€¢" for
#: a bullet, "Ã©" for é. Some official sites publish it that way. A run is a
#: UTF-8 lead byte followed by its continuation bytes, each seen through the
#: Windows-1252 table; only such runs are touched, so genuine accented text
#: beside them is left exactly as it is.
_CONTINUATION = "\u0080-\u00bf\u20ac\u201a\u0192\u201e\u2026\u2020\u2021\u02c6\u2030\u0160\u2039\u0152\u017d\u2018\u2019\u201c\u201d\u2022\u2013\u2014\u02dc\u2122\u0161\u203a\u0153\u017e\u0178"
MOJIBAKE = re.compile(f"[\u00c2-\u00df][{_CONTINUATION}]|[\u00e0-\u00ef][{_CONTINUATION}]{{2}}|[\u00f0-\u00f4][{_CONTINUATION}]{{3}}")


def _repair_run(match: re.Match) -> str:
    run = match.group(0)
    try:
        return run.encode("cp1252").decode("utf-8")
    except (UnicodeEncodeError, UnicodeDecodeError):
        return run


def clean_text(text: str) -> str:
    """Prose with the invisible characters removed, stray non-breaking spaces made
    plain, and double-encoded runs put back the way their author typed them."""
    text = ZERO_WIDTH.sub("", text or "")
    text = MOJIBAKE.sub(_repair_run, text)
    return text.replace("\u00a0", " ")

_splitter = SentenceSplitter(chunk_size=CHUNK_TOKENS, chunk_overlap=CHUNK_OVERLAP)


class DocumentUnreadable(ValueError):
    """Nothing quotable could be read out of the file.

    Raised rather than returned empty because the caller's next act is to
    supersede the previous version of this document. An upload that yields
    nothing must never be allowed to retire knowledge that works: the document
    would go mute, the upload would report success, and nobody would find out
    until an answer that used to cite it stopped doing so.
    """


def extension_of(filename: str) -> str:
    name = (filename or "").lower()
    dot = name.rfind(".")
    return name[dot:] if dot != -1 else ""


def is_supported(filename: str) -> bool:
    return extension_of(filename) in SUPPORTED_EXTENSIONS


# --------------------------------------------------------------------------
# Readers. Each returns plain text; the splitter turns text into passages.
# --------------------------------------------------------------------------


def _text_from_docx(data: bytes) -> str:
    """Paragraphs and tables, in the order they appear in the document body.

    Tables are walked explicitly because python-docx keeps their text out of
    `.paragraphs`, and an official circular puts its dates and figures in a
    table as often as in a sentence.
    """
    from docx import Document as DocxDocument
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = DocxDocument(io.BytesIO(data))
    out: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            # A heading becomes a markdown heading so the section parser can
            # see it, which is what gives a fragment its heading path.
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(c for c in style if c.isdigit()) or "1"
                out.append(f"{'#' * min(int(level), 6)} {text}")
            else:
                out.append(text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    out.append(" | ".join(cells))
    return "\n\n".join(out)


def _text_from_odt(data: bytes) -> str:
    from odf import teletype
    from odf.opendocument import load

    document = load(io.BytesIO(data))
    parts: list[str] = []
    for element in document.text.childNodes:
        rendered = teletype.extractText(element).strip()
        if not rendered:
            continue
        if element.qname[1] == "h":
            level = element.getAttribute("outlinelevel") or "1"
            parts.append(f"{'#' * min(int(level), 6)} {rendered}")
        else:
            parts.append(rendered)
    return "\n\n".join(parts)


def _text_from_rtf(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    return rtf_to_text(data.decode("utf-8", errors="replace"), errors="ignore")


def _normalise_pdf_text(text: str) -> str:
    """Undo the damage layout extraction does to ordinary prose.

    Layout mode preserves the visual arrangement, which is what lets headings
    and tables survive, but it also pads with runs of spaces and breaks lines
    mid-sentence. Left in, those runs reach the embedding as noise and the
    quote shown to a visitor looks broken.
    """
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t]{2,}", "  ", text)
    # A line ending mid-sentence is a wrap, not a break.
    text = re.sub(r"(?<![.!?:;])\n(?=[a-z(])", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _promote_allcaps_headings(text: str) -> str:
    """SECTION TITLES IN CAPS become headings the section parser can see.

    Official documents rarely carry style information, and a capitalised line
    on its own is how they mark a section. Recognising it is the difference
    between a fragment that knows it came from "PART III, COMMENCEMENT" and
    one that does not.
    """
    out: list[str] = []
    for line in text.split("\n"):
        stripped = line.strip()
        letters = [c for c in stripped if c.isalpha()]
        looks_like_heading = (
            3 <= len(stripped) <= 90
            and len(letters) >= 3
            and all(c.isupper() for c in letters)
            and not stripped.endswith((".", ",", ";"))
        )
        out.append(f"## {stripped}" if looks_like_heading else line)
    return "\n".join(out)


def _pdf_pages(data: bytes) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return [page.extract_text(extraction_mode="layout") or "" for page in reader.pages]


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------


def _passages(text: str) -> list[tuple[str, str]]:
    """Split on sections first, then on sentences, as (text, heading path).

    Section first is what keeps a heading attached. Splitting straight to
    sentences would give passages that read fine and cite nothing.

    The section parser reports only a section's ANCESTORS; the section's own
    heading is the first line of its content. Both are needed, or every
    passage is filed one level above where it actually sat.
    """
    text = clean_text(text)
    sections = MarkdownNodeParser().get_nodes_from_documents([LlamaDocument(text=text)])
    out: list[tuple[str, str]] = []
    for section in sections:
        body = section.get_content().strip()
        if not body:
            continue
        ancestors = [p.strip() for p in str(section.metadata.get("header_path", "")).split("/") if p.strip()]
        own = ""
        first = body.split("\n", 1)[0].strip()
        if first.startswith("#"):
            own = first.lstrip("#").strip()
        path = " › ".join(dict.fromkeys([*ancestors, own] if own else ancestors))[:500]
        # The heading stays in the passage as well as in its path: a fragment
        # that carries its own heading words is found by the words people
        # search for it with.
        for child in _splitter.get_nodes_from_documents([LlamaDocument(text=body)]):
            piece = child.get_content().strip()
            if piece:
                out.append((piece, path))
    return out


def build_chunks(filename: str, data: bytes) -> list[dict]:
    """Parse one document into passages: {"text", "page", "heading_path"}.

    Raises DocumentUnreadable for anything that yields nothing quotable, which
    includes the common and important case of a scanned PDF with no text layer.
    """
    extension = extension_of(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentUnreadable(
            f"{extension or 'That file'} is not a format Ma’at reads. "
            f"Accepted: {', '.join(e[1:].upper() for e in SUPPORTED_EXTENSIONS)}."
        )

    try:
        chunks: list[dict] = []
        if extension == ".pdf":
            pages = _pdf_pages(data)
            for number, raw in enumerate(pages, start=1):
                page_text = _promote_allcaps_headings(_normalise_pdf_text(raw))
                if not page_text.strip():
                    continue
                for body, path in _passages(page_text):
                    chunks.append({"text": body, "page": number, "heading_path": path})
        else:
            if extension == ".docx":
                text = _text_from_docx(data)
            elif extension == ".odt":
                text = _text_from_odt(data)
            elif extension == ".rtf":
                text = _text_from_rtf(data)
            else:
                text = data.decode("utf-8", errors="replace")
            text = _promote_allcaps_headings(text)
            for body, path in _passages(text):
                chunks.append({"text": body, "page": None, "heading_path": path})
    except DocumentUnreadable:
        raise
    except Exception as exc:  # noqa: BLE001 - one boundary, on purpose
        raise DocumentUnreadable(
            f"This file could not be read, and may be corrupt or not really "
            f"a {extension[1:].upper()} file ({exc.__class__.__name__})."
        ) from exc

    if not chunks:
        raise DocumentUnreadable(_nothing_read(extension))
    return chunks


def _nothing_read(extension: str) -> str:
    if extension == ".pdf":
        return (
            "No text could be read from this PDF. It is most likely a scan, "
            "which is an image of a page rather than the words on it. Upload a "
            "text PDF, or the original document it was made from."
        )
    return "This file has no readable text in it."
